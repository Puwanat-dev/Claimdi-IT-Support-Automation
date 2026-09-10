import os
import openpyxl
import pyperclip
import tkinter as tk
from tkinter import messagebox

class Mail_Answer:
    DEFAULT_PATTERN_FILE = os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "mail_patterns.xlsx" #file name
    )
    
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Mail Answer Pattern Selector")
        self.root.geometry("500x600")
        
        # Read column data
        self.A_patterns = self.read_Column_A()
        self.B_patterns = self.read_Column_B()
        
        # Create UI
        
        self.create_A_buttons() 
        self.create_B_buttons()
        

    def read_Column_A(self):
        """Read all cells from column A until empty cell is found"""
        A_patterns = []
        
        try:
            # Check if file exists
            if not os.path.exists(self.DEFAULT_PATTERN_FILE):
                messagebox.showwarning("File Not Found", 
                                     f"Pattern file not found at:\n{self.DEFAULT_PATTERN_FILE}")
                return A_patterns
            
            # Open workbook and get active sheet
            workbook = openpyxl.load_workbook(self.DEFAULT_PATTERN_FILE)
            sheet = workbook.active
            
            # Read column A starting from row 1
            row = 1
            while True:
                cell_value = sheet.cell(row=row, column=1).value
                
                # Break when empty cell is found
                if cell_value is None or str(cell_value).strip() == "":
                    break
                    
                A_patterns.append(str(cell_value).strip())
                row += 1
            
            workbook.close()
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read pattern file:\n{str(e)}")
        
        return A_patterns
    
    def create_A_buttons(self):
        """Create buttons for each pattern found in column A"""
        if not self.A_patterns:
            label = tk.Label(self.root, text="No patterns found in file", 
                           font=("Arial", 12), fg="red")
            label.pack(pady=20)
            return
        
        # Title label
        title_label = tk.Label(self.root, text="Select Mail Pattern:", 
                              font=("Arial", 14, "bold"))
        title_label.pack(pady=10)
        
        # Create a scrollable frame for buttons
        frame_canvas = tk.Frame(self.root)
        frame_canvas.pack(fill="both", expand=True, padx=20, pady=10)
        
        canvas = tk.Canvas(frame_canvas)
        scrollbar = tk.Scrollbar(frame_canvas, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas)
        
        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        
        # Create buttons for each pattern
        for i, pattern in enumerate(self.A_patterns):
            button = tk.Button(
                scrollable_frame, 
                text=pattern,
                command=lambda p=pattern: self.copy_to_clipboard(p),
                width=35,
                height=2,
                bg="#f0f0f0",
                relief="raised",
                cursor="hand2"
            )
            button.pack(pady=5)
        
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")



    def read_Column_B(self):
        """Read all cells from column B until empty cell is found"""
        B_patterns = []
        
        try:
            # Check if file exists
            if not os.path.exists(self.DEFAULT_PATTERN_FILE):
                messagebox.showwarning("File Not Found", 
                                     f"Pattern file not found at:\n{self.DEFAULT_PATTERN_FILE}")
                return B_patterns
            
            # Open workbook and get active sheet
            workbook = openpyxl.load_workbook(self.DEFAULT_PATTERN_FILE)
            sheet = workbook.active
            
            # Read column B starting from row 1
            row = 1
            while True:
                cell_value = sheet.cell(row=row, column=2).value
                
                # Break when empty cell is found
                if cell_value is None or str(cell_value).strip() == "":
                    break
                    
                B_patterns.append(str(cell_value).strip())
                row += 1
            
            workbook.close()
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read pattern file:\n{str(e)}")
        
        return B_patterns    

    
    def create_B_buttons(self):
        """Create buttons for each pattern found in column B"""
        if not self.B_patterns:
            label = tk.Label(self.root, text="No patterns found in file", 
                           font=("Arial", 12), fg="red")
            label.pack(pady=20)
            return
        
        # Title label
        title_label = tk.Label(self.root, text="Select Mail Pattern:", 
                              font=("Arial", 14, "bold"))
        title_label.pack(pady=10)
        
        # Create a scrollable frame for buttons
        frame_canvas = tk.Frame(self.root)
        frame_canvas.pack(fill="both", expand=True, padx=20, pady=10)
        
        canvas = tk.Canvas(frame_canvas)
        scrollbar = tk.Scrollbar(frame_canvas, orient="vertical", command=canvas.yview)
        scrollable_frame = tk.Frame(canvas)
        
        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )
        
        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        
        # Create buttons for each pattern
        for i, pattern in enumerate(self.B_patterns):
            button = tk.Button(
                scrollable_frame, 
                text=pattern,
                command=lambda p=pattern: self.copy_to_clipboard(p),
                width=35,
                height=2,
                bg="#f0f0f0",
                relief="raised",
                cursor="hand2"
            )
            button.pack(pady=5)
        
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
    
    
    def copy_to_clipboard(self, text):
        """Copy selected pattern to clipboard"""
        pyperclip.copy(text)
        messagebox.showinfo("Success", f"Copied to clipboard:\n{text}")
    
    def run(self):
        """Start the Tkinter main loop"""
        self.root.mainloop()


if __name__ == "__main__":
    app = Mail_Answer()
    app.run()