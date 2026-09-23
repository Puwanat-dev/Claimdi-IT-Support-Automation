import os
import openpyxl
import pyperclip
import tkinter as tk
from tkinter import messagebox
from read_config import ReadConfig

class Mail_Answer:
    DEFAULT_PATTERN_FILE = os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "mail_patterns.xlsx" #file name
    )
    
    def __init__(self):

        self.root = tk.Tk()
        self.root.title("Mail Answer Pattern Selector")
        self.root.geometry("900x620")
        self.root.configure(bg="#f5f5f5")

        #set variable for mail patterns and config
        self.mail_patterns = []
        self.config = ReadConfig("config.txt")
        self.it_support_name = self.config.read_variable("it_support_name")
        self.it_support_phone = self.config.read_variable("it_support_phone")
       
        self.status_var = tk.StringVar(value="No text appended")

        # Read column data
        self.A_patterns = self.read_Column_A()
        self.B_patterns = self.read_Column_B()

        # Create UI
        main_frame = tk.Frame(self.root, bg="#f5f5f5")
        main_frame.pack(fill="both", expand=True, padx=20, pady=20)

        top_frame = tk.Frame(main_frame, bg="#f5f5f5")
        top_frame.pack(fill="x", pady=(0, 10))

        action_row = tk.Frame(top_frame, bg="#f5f5f5")
        action_row.pack(fill="x")
        self.create_copy_buttons(action_row)
        self.create_reset_button(action_row)

        status_bar = tk.Label(
            top_frame,
            textvariable=self.status_var,
            bg="#ffffff",
            fg="#222222",
            relief="sunken",
            bd=1,
            anchor="w",
            justify="left",
            padx=8,
            pady=6,
            font=("Arial", 10),
        )
        status_bar.pack(fill="x", pady=(10, 0))

        columns_frame = tk.Frame(main_frame, bg="#f5f5f5")
        columns_frame.pack(fill="both", expand=True)
        columns_frame.grid_columnconfigure(0, weight=1)
        columns_frame.grid_columnconfigure(1, weight=1)

        self.create_A_buttons(columns_frame)
        self.create_B_buttons(columns_frame)

        

    def update_status(self):
        if not self.mail_patterns:
            self.status_var.set("No text appended")
        else:
            self.status_var.set("Appended: " + " | ".join(self.mail_patterns))

    def create_copy_buttons(self, parent):
        copy_button = tk.Button(
            parent,
            text="Copy to Clipboard",


            command=lambda: 

            #Mail script format
            self.copy_to_clipboard(f"เรียน {self.mail_patterns[0]} \n\n{self.mail_patterns[1]} \n\nBest Regards,\n{self.it_support_name}\nIT Support\nAnywhere 2 go Co.,Ltd.\n520 Dindaeng Rd., Bangkok 10400, Thailand\nTel: {self.it_support_phone}"),

            width=20,
            height=2,
            bg="#f0f0f0",
            relief="raised",
            cursor="hand2",
        )
        copy_button.pack(side="left", padx=(0, 10), pady=5)

    def create_reset_button(self, parent):
        reset_button = tk.Button(
            parent,
            text="Reset",
            command=lambda: (self.mail_patterns.clear(), self.update_status()),
            width=15,
            height=2,
            bg="#f0f0f0",
            relief="raised",
            cursor="hand2",
        )
        reset_button.pack(side="left", pady=5)


    def read_Column_A(self):
        """Read all cells from column A until empty cell is found"""
        A_patterns = []
        
        try:
            # Check if file exists
            if not os.path.exists(self.DEFAULT_PATTERN_FILE):
                messagebox.showwarning("File Not Found", 
                                     f"Pattern file not found at:\n{self.DEFAULT_PATTERN_FILE}")
                return A_patterns
            
            # Open workbook and get active sheet
            workbook = openpyxl.load_workbook(self.DEFAULT_PATTERN_FILE)
            sheet = workbook.active
            
            # Read column A starting from row 1
            row = 1
            while True:
                cell_value = sheet.cell(row=row, column=1).value
                
                # Break when empty cell is found
                if cell_value is None or str(cell_value).strip() == "":
                    break
                    
                A_patterns.append(str(cell_value).strip())
                row += 1
            
            workbook.close()
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read pattern file:\n{str(e)}")
        
        return A_patterns
    
    def create_A_buttons(self, parent):
        """Create buttons for each pattern found in column A"""
        if not self.A_patterns:
            label = tk.Label(parent, text="No A patterns found in file",
                           font=("Arial", 12), fg="red", bg="#f5f5f5")
            label.grid(row=0, column=0, padx=10, pady=10, sticky="n")
            return

        section = tk.Frame(parent, bg="#f5f5f5")
        section.grid(row=0, column=0, padx=(0, 10), pady=10, sticky="nsew")

        title_label = tk.Label(section, text="A Patterns", font=("Arial", 14, "bold"), bg="#f5f5f5")
        title_label.pack(pady=(0, 10))

        frame_canvas = tk.Frame(section, bg="#f5f5f5")
        frame_canvas.pack(fill="both", expand=True)

        canvas = tk.Canvas(frame_canvas, bg="#f5f5f5")
        scrollbar = tk.Scrollbar(frame_canvas, orient="vertical", command=canvas.yview , width=30)
        scrollable_frame = tk.Frame(canvas, bg="#f5f5f5")

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        for pattern in self.A_patterns:
            button = tk.Button(
                scrollable_frame,
                text=pattern,
                command=lambda p=pattern: (self.mail_patterns.append(p), self.update_status()),
                width=35,
                height=2,
                bg="#f0f0f0",
                relief="raised",
                cursor="hand2",
            )
            button.pack(pady=5, fill="x")

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")



    def read_Column_B(self):
        """Read all cells from column B until empty cell is found"""
        B_patterns = []
        
        try:
            # Check if file exists
            if not os.path.exists(self.DEFAULT_PATTERN_FILE):
                messagebox.showwarning("File Not Found", 
                                     f"Pattern file not found at:\n{self.DEFAULT_PATTERN_FILE}")
                return B_patterns
            
            # Open workbook and get active sheet
            workbook = openpyxl.load_workbook(self.DEFAULT_PATTERN_FILE)
            sheet = workbook.active
            
            # Read column B starting from row 1
            row = 1
            while True:
                cell_value = sheet.cell(row=row, column=2).value
                
                # Break when empty cell is found
                if cell_value is None or str(cell_value).strip() == "":
                    break
                    
                B_patterns.append(str(cell_value).strip())
                row += 1
            
            workbook.close()
            
        except Exception as e:
            messagebox.showerror("Error", f"Failed to read pattern file:\n{str(e)}")
        
        return B_patterns    

    
    def create_B_buttons(self, parent):
        """Create buttons for each pattern found in column B"""
        if not self.B_patterns:
            label = tk.Label(parent, text="No B patterns found in file",
                           font=("Arial", 12), fg="red", bg="#f5f5f5")
            label.grid(row=0, column=1, padx=10, pady=10, sticky="n")
            return

        section = tk.Frame(parent, bg="#f5f5f5")
        section.grid(row=0, column=1, padx=(10, 0), pady=10, sticky="nsew")

        title_label = tk.Label(section, text="B Patterns", font=("Arial", 14, "bold"), bg="#f5f5f5")
        title_label.pack(pady=(0, 10))

        frame_canvas = tk.Frame(section, bg="#f5f5f5")
        frame_canvas.pack(fill="both", expand=True)

        canvas = tk.Canvas(frame_canvas, bg="#f5f5f5")
        scrollbar = tk.Scrollbar(frame_canvas, orient="vertical", command=canvas.yview, width=30)
        scrollable_frame = tk.Frame(canvas, bg="#f5f5f5")

        scrollable_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all"))
        )

        canvas.create_window((0, 0), window=scrollable_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)

        for pattern in self.B_patterns:
            button = tk.Button(
                scrollable_frame,
                text=pattern,
                command=lambda p=pattern: (self.mail_patterns.append(p), self.update_status()),
                width=50,
                height=2,
                bg="#f0f0f0",
                relief="raised",
                cursor="hand2",
            )
            button.pack(pady=5, fill="x")

        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
    
    
    def copy_to_clipboard(self, text):
        """Copy selected pattern to clipboard"""
        pyperclip.copy(text)
        messagebox.showinfo("Success", f"Copied to clipboard:\n{text}")
    
    def run(self):
        """Start the Tkinter main loop"""
        self.root.mainloop()


if __name__ == "__main__":
    app = Mail_Answer()
    app.run()