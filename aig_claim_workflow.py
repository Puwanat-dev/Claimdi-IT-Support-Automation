"""
AIG workflow:
  1. Pick an .xlsx file via a file dialog
  2. Count rows with data in Column A (excluding row 1)
  3. In Column I, replace empty values with a single space
  4. Delete Column T
  5. Append a report (filename, row count, timestamp) to ~/Desktop/report.txt
  6. Save a copy next to the original, with the first 4 characters of the
     filename (excluding the extension) removed

Requires: pip install openpyxl
"""

from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog

from openpyxl import load_workbook


class AIGWorkflow:
    COL_A = 1      # Column A
    COL_I = 9      # Column I
    COL_T = 20     # Column T
    CHARS_TO_STRIP = 4

    def __init__(self):
        self.file_path: Path | None = None
        self.workbook = None
        self.sheet = None
        self.row_count: int = 0
        self.output_path: Path | None = None

    # ------------------------------------------------------------------ 1
    def select_file(self) -> Path:
        """Open a file dialog and let the user choose an .xlsx file."""
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)  # keep the dialog in front
        selected = filedialog.askopenfilename(
            title="Select an Excel file",
            filetypes=[("Excel files", "*.xlsx")],
        )
        root.destroy()

        if not selected:
            raise SystemExit("No file selected. Exiting.")

        self.file_path = Path(selected)
        return self.file_path

    def load_workbook(self) -> None:
        self.workbook = load_workbook(self.file_path)
        self.sheet = self.workbook.active

    # ------------------------------------------------------------------ 2
    @staticmethod
    def _is_empty(value) -> bool:
        return value is None or (isinstance(value, str) and value.strip() == "")

    def count_rows(self) -> int:
        """Count rows (excluding row 1) where Column A contains data."""
        count = 0
        for (value,) in self.sheet.iter_rows(
            min_row=2, min_col=self.COL_A, max_col=self.COL_A, values_only=True
        ):
            if not self._is_empty(value):
                count += 1
        self.row_count = count
        return count

    # ------------------------------------------------------------------ 3
    def fill_empty_in_column_i(self) -> None:
        """Replace empty values in Column I with a single space (data rows only)."""
        for row in range(2, self.sheet.max_row + 1):
            cell = self.sheet.cell(row=row, column=self.COL_I)
            if cell.value is None or cell.value == "":
                cell.value = " "

    # ------------------------------------------------------------------ 4
    def delete_column_t(self) -> None:
        """Delete Column T (columns to the right shift left)."""
        self.sheet.delete_cols(self.COL_T)

    # ------------------------------------------------------------------ 5
    def save_report(self) -> Path:
        """Append a report to report.txt on the Desktop (created if missing)."""
        desktop = Path.home() / "Desktop"
        desktop.mkdir(parents=True, exist_ok=True)
        report_path = desktop / "report.txt"

        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        lines = (
            f"Original filename: {self.file_path.name}\n"
            f"Row count (Column A, excluding row 1): {self.row_count}\n"
            f"Timestamp: {timestamp}\n"
            f"{'-' * 40}\n"
        )
        with open(report_path, "a", encoding="utf-8") as f:
            f.write(lines)
        return report_path

    # ------------------------------------------------------------------ 6
    def save_modified_workbook(self) -> Path:
        """Save as a new file in the same folder, first 4 name characters removed."""
        stem = self.file_path.stem
        if len(stem) <= self.CHARS_TO_STRIP:
            raise ValueError(
                f"Filename '{stem}' has {len(stem)} characters; "
                f"it must be longer than {self.CHARS_TO_STRIP} to remove the first "
                f"{self.CHARS_TO_STRIP}."
            )

        new_name = stem[self.CHARS_TO_STRIP:] + self.file_path.suffix
        self.output_path = self.file_path.with_name(new_name)
        self.workbook.save(self.output_path)
        return self.output_path

    # ---------------------------------------------------------------- run
    def run(self) -> None:
        self.select_file()
        self.load_workbook()
        self.count_rows()
        self.fill_empty_in_column_i()
        self.delete_column_t()
        report_path = self.save_report()
        saved_path = self.save_modified_workbook()

        print(f"Rows with data in Column A: {self.row_count}")
        print(f"Report written to: {report_path}")
        print(f"Modified workbook saved to: {saved_path}")


if __name__ == "__main__":
    AIGWorkflow().run()
